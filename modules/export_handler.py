from fpdf import FPDF
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
import os
import re
from config import Config
from datetime import datetime


class ExportHandler:

    @staticmethod
    def _safe_pdf_text(text):
        """FPDF-safe encoding wrapper to prevent Unicode crashes."""
        if not text: return ""
        if not isinstance(text, str): text = str(text)
        
        # Mapping common problematic Unicode to safe Latin-1
        replacements = {
            '\u201c': '"', '\u201d': '"',  # Smart quotes
            '\u2018': "'", '\u2019': "'",  # Smart apostrophes
            '\u2014': '-', '\u2013': '-',  # Em/En dashes
            '\u2026': '...',               # Ellipsis
            '\u2212': '-',                # Minus sign
            '\u20b9': 'Rs.',              # Rupee symbol
        }
        for u_char, r_char in replacements.items():
            text = text.replace(u_char, r_char)
            
        # Final fallback for anything else
        try:
            return text.encode('latin-1', 'replace').decode('latin-1')
        except:
            return text.encode('ascii', 'replace').decode('ascii')

    @staticmethod
    def _clean_text(text):
        """Strip markdown and leftover metadata tags for clean export."""
        if not text:
            return ""
        # 1. Aggressively remove all bracketed metadata like [Short Answer], [Conceptual], [Application]
        # This matches anything inside brackets and trailing whitespace
        text = re.sub(r'\[[^\]]+\]\s*', '', text)
        
        # 2. Standard markdown cleaning
        text = re.sub(r'```\w*\n?', '', text)
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'`([^`]+)`', r'\1', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    # ═══════════════════════════════════════════════════════════════════════════
    # PDF EXPORT
    # ═══════════════════════════════════════════════════════════════════════════

    @staticmethod
    def export_to_pdf(paper, filename=None):
        """Export question paper to a clean, professional PDF."""

        if not filename:
            filename = f"{paper['name'].replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        filepath = os.path.join(Config.OUTPUT_DIR, filename)

        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=20)

        # ── Grouping & Sorting (MCQ FIRST) ──
        type_priority = {"MCQ": 0, "Fill in the Blanks": 1, "Short Answer": 2, "Long Answer": 3}
        questions = sorted(paper["questions"], key=lambda x: (type_priority.get(x.get("type"), 4), x.get("no", 0)))

        # ── Title (Professional Layout) ──
        pdf.set_font('Times', 'B', 22)
        exam_name = ExportHandler._safe_pdf_text(paper.get('exam_name', 'EXAMINATION QUESTION PAPER'))
        pdf.cell(0, 12, exam_name.upper(), 0, 1, 'C')
        
        # Course Detail
        pdf.set_font('Times', 'B', 15)
        course_str = f"Course: {paper.get('course_name','')} ({paper.get('course_code','')})"
        pdf.cell(0, 10, ExportHandler._safe_pdf_text(course_str), 0, 1, 'C')
        
        # Meta Details Line
        pdf.set_font('Times', '', 11)
        duration = paper.get('duration', '3 Hours')
        max_marks = paper.get('max_marks', paper.get('total_marks', 0))
        date_str = datetime.now().strftime('%d-%m-%Y')
        meta_line = f"Set: {paper.get('set','A')}   |   Time: {duration}   |   Max Marks: {max_marks}   |   Date: {date_str}"
        pdf.cell(0, 8, ExportHandler._safe_pdf_text(meta_line), 0, 1, 'C')
        pdf.ln(4)

        # ── Student Header ──
        pdf.set_font('Times', 'B', 11)
        pdf.cell(95, 8, "Name: __________________________", 0, 0)
        pdf.cell(0, 8, "Roll No: _______________________", 0, 1, 'R')
        pdf.cell(95, 8, "Student Sign: __________________", 0, 0)
        pdf.cell(0, 8, "Teacher Sign: __________________", 0, 1, 'R')
        pdf.ln(2)
        pdf.set_draw_color(0, 0, 0)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(6)

        # ── Instructions Box (Professional) ──
        instr = paper.get('config', {}).get('instructions', '1. Answer all questions.\n2. Figures to the right indicate full marks.')
        pdf.set_fill_color(250, 250, 250)
        pdf.set_font('Times', 'B', 10)
        pdf.cell(0, 7, "GENERAL INSTRUCTIONS:", 1, 1, 'L', True)
        pdf.set_font('Times', '', 10)
        pdf.multi_cell(0, 6, ExportHandler._safe_pdf_text(instr), 1, 'L')
        pdf.ln(8)

        # ── Render Sections and Questions ──
        sections = ["A", "B", "C", "D", "E"]
        current_type = None
        type_idx = 0
        q_count = 0

        for q in questions:
            # Add Section Header when type changes (Left Aligned, No Box)
            if q.get("type") != current_type:
                current_type = q.get("type")
                sec_letter = sections[type_idx] if type_idx < len(sections) else chr(65 + type_idx)
                
                pdf.ln(10)
                pdf.set_font('Times', 'B', 14)
                header_text = f"SECTION {sec_letter}: {current_type.upper()}S"
                if "MCQ" in current_type.upper(): header_text = f"SECTION {sec_letter}: MCQs"
                pdf.cell(0, 8, header_text, 0, 1, 'L')
                pdf.ln(4)
                type_idx += 1
            
            q_count += 1
            
            # Question and Marks line
            pdf.set_font('Times', 'B', 11)
            pdf.write(6, f"{q_count}. ")
            
            pdf.set_font('Times', '', 11)
            content = ExportHandler._clean_text(q.get('q', q.get('content', '')))
            content = ExportHandler._safe_pdf_text(content)
            pdf.write(6, content)
            
            pdf.set_font('Times', 'B', 11)
            pdf.cell(0, 6, f" [{q['marks']} Marks]", 0, 1, 'R')
            
            # ── MCQ Options ──
            opts = q.get('options', [])
            if q.get("type") == "MCQ" and opts and isinstance(opts, list):
                pdf.ln(2)
                pdf.set_font('Times', '', 10)
                prefixes = ['a)', 'b)', 'c)', 'd)']
                for i, opt in enumerate(opts[:4]):
                    pdf.set_x(25)
                    opt_txt = ExportHandler._safe_pdf_text(str(opt))
                    pdf.cell(0, 6, f"{prefixes[i]} {opt_txt}", 0, 1)
            
            pdf.ln(8)

        # ── End ──
        pdf.ln(5)
        pdf.set_draw_color(100, 100, 100)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(3)
        pdf.set_font('Arial', 'B', 10)
        pdf.cell(0, 8, '--- End of Question Paper ---', 0, 1, 'C')

        pdf.output(filepath)
        return filepath

    # ═══════════════════════════════════════════════════════════════════════════
    # DOCX EXPORT
    # ═══════════════════════════════════════════════════════════════════════════

    @staticmethod
    def export_to_docx(paper, filename=None):
        """Export question paper to a professional DOCX."""

        if not filename:
            filename = f"{paper['name'].replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = os.path.join(Config.OUTPUT_DIR, filename)

        doc = Document()

        # ── Global Typography (Academic) ──
        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(11)

        # ── Title (Professional Layout) ──
        exam_name = paper.get('exam_name', paper.get('title', 'EXAMINATION')).upper()
        tp = doc.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tr = tp.add_run(exam_name)
        tr.bold = True
        tr.font.size = Pt(22)

        # Course Detail
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(f"Course: {paper.get('course_name','')} ({paper.get('course_code','')})")
        cr.bold = True
        cr.font.size = Pt(16)

        # Meta Details Line
        max_marks = paper.get('max_marks', paper.get('total_marks', 0))
        date_str = datetime.now().strftime('%d-%m-%Y')
        dp = doc.add_paragraph()
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_text = f"Set: {paper.get('set','A')}  |  Time: {paper.get('duration','3 Hours')}  |  Max Marks: {max_marks}  |  Date: {date_str}"
        dr = dp.add_run(meta_text)
        dr.font.size = Pt(11)

        doc.add_paragraph('_' * 80).alignment = WD_ALIGN_PARAGRAPH.CENTER

        # ── Instructions Box (Professional) ──
        instr = paper.get('config', {}).get('instructions', '1. Answer all questions.\n2. Figures to the right indicate full marks.')
        inst_table = doc.add_table(rows=1, cols=1)
        inst_table.style = 'Table Grid'
        cell = inst_table.rows[0].cells[0]
        
        # Heading
        iph = cell.paragraphs[0]
        ipr = iph.add_run("GENERAL INSTRUCTIONS:")
        ipr.bold = True
        ipr.font.size = Pt(10)
        
        # Content
        ipc = cell.add_paragraph(instr)
        ipc.paragraph_format.space_before = Pt(4)
        if ipc.runs:
            ipc.runs[0].font.size = Pt(10)
        
        doc.add_paragraph() # Spacer

        # ── Grouping & Sorting (MCQ FIRST) ──
        type_priority = {"MCQ": 0, "Fill in the Blanks": 1, "Short Answer": 2, "Long Answer": 3}
        questions = sorted(paper["questions"], key=lambda x: (type_priority.get(x.get("type"), 4), x.get("no", 0)))

        # ── Header Fields ──
        h_table = doc.add_table(rows=2, cols=2)
        h_table.width = Inches(6.5)
        h_table.cell(0, 0).text = "Name: ____________________"
        h_table.cell(0, 1).text = "Roll No: ________________"
        h_table.cell(1, 0).text = "Student Sign: ____________"
        h_table.cell(1, 1).text = "Teacher Sign: ____________"
        for row in h_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.runs[0].font.size = Pt(11)
                    paragraph.runs[0].font.bold = True

        doc.add_paragraph()

        # ── Render Sections and Questions ──
        sections = ["A", "B", "C", "D", "E"]
        current_type = None
        type_idx = 0
        q_count = 0

        for q in questions:
            # Section Header (Left Aligned, Bold)
            if q.get("type") != current_type:
                current_type = q.get("type")
                sec_letter = sections[type_idx] if type_idx < len(sections) else chr(65 + type_idx)
                type_idx += 1
                
                sh = doc.add_paragraph()
                sh.paragraph_format.space_before = Pt(18)
                header_text = f"SECTION {sec_letter}: {current_type.upper()}S"
                if "MCQ" in current_type.upper(): header_text = f"SECTION {sec_letter}: MCQs"
                sh_run = sh.add_run(header_text)
                sh_run.bold = True
                sh_run.font.size = Pt(12)

            q_count += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            
            # Question Index
            p.add_run(f"{q_count}.  ").bold = True

            # Marks on Right: [X Marks] using Tab Stop
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(f"\t[{q['marks']} Marks]").bold = True

            # Question Text (Cleaned)
            content = ExportHandler._clean_text(q.get('q', q.get('content', '')))
            cp = doc.add_paragraph(content)
            cp.paragraph_format.left_indent = Inches(0.3)
            cp.paragraph_format.right_indent = Inches(0.8) # Keep space for marks
            cp.paragraph_format.space_after = Pt(6)

            # MCQ Options (Vertical)
            opts = q.get('options', [])
            if q.get("type") == "MCQ" and opts and isinstance(opts, list):
                prefixes = ['a)', 'b)', 'c)', 'd)']
                for i, opt in enumerate(opts[:4]):
                    op = doc.add_paragraph()
                    op.paragraph_format.left_indent = Inches(0.6)
                    op.add_run(f"{prefixes[i]} {opt}").font.size = Pt(10)
            
            doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # ── End ──
        doc.add_paragraph('_' * 80)
        end = doc.add_paragraph('--- End of Question Paper ---')
        end.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if end.runs:
            end.runs[0].bold = True

        doc.save(filepath)
        return filepath

    # ═══════════════════════════════════════════════════════════════════════════
    # TXT EXPORT
    # ═══════════════════════════════════════════════════════════════════════════

    @staticmethod
    def _to_txt(paper):
        """User's refined text formatter with absolute safety."""
        cfg = paper.get("config", {})
        exam_name = paper.get('exam_name', paper.get('title', 'Examination')).upper()
        
        lines = []
        lines.append("="*65)
        lines.append(f"{exam_name:^65}")
        
        c_name = paper.get("course_name", "")
        c_code = paper.get("course_code", "")
        if c_name:
            lines.append(f"Course: {c_name} ({c_code})".center(65))
        
        duration = paper.get('duration', '3 Hours')
        max_marks = paper.get('max_marks', paper.get('total_marks', 0))
        date_str = datetime.now().strftime('%d-%m-%Y')
        meta_text = f"Set: {paper.get('set','A')}  |  Time: {duration}  |  Max Marks: {max_marks}  |  Date: {date_str}"
        lines.append(meta_text.center(65))
        lines.append("="*65)
        
        # General Instructions
        instr = cfg.get('instructions', '1. Answer all questions.\n2. Figures to the right indicate full marks.')
        lines.append("\nGENERAL INSTRUCTIONS:")
        lines.append("-" * 21)
        lines.append(instr)
        lines.append("\n" + "_"*65 + "\n")
        
        lines.append(f"Name: ____________________    Roll No: ________________")
        lines.append(f"Student Sign: ____________    Teacher Sign: ____________")
        lines.append(f"{'═'*65}")

        # ── Grouping & Sorting (MCQ FIRST) ──
        type_priority = {"MCQ": 0, "Fill in the Blanks": 1, "Short Answer": 2, "Long Answer": 3}
        questions = sorted(paper.get("questions", []), key=lambda x: (type_priority.get(x.get("type"), 4), x.get("no", 0)))

        sections = ["A", "B", "C", "D", "E"]
        current_type = None
        type_idx = 0
        q_count = 0

        # ── Render Questions (Sequential 1, 2, 3...) ──
        for q in questions:
            # Section Header (Left Aligned)
            if q.get("type") != current_type:
                current_type = q.get("type")
                sec_letter = sections[type_idx] if type_idx < len(sections) else chr(65 + type_idx)
                header_text = f"SECTION {sec_letter}: {current_type.upper()}S"
                if "MCQ" in current_type.upper(): header_text = f"SECTION {sec_letter}: MCQs"
                lines.append(f"\n{header_text}")
                lines.append(f"{'─'*len(header_text)}\n")
                type_idx += 1

            q_count += 1
            q_txt = ExportHandler._clean_text(q.get('q', q.get('content', 'No question text')))
            
            # Question Index + Marks
            lines.append(f"{q_count}. {q_txt}")
            lines.append(f"{' '*(60 - len(str(q.get('marks', 0))))} [{q.get('marks', 0)} Marks]")
            
            # ── MCQ Options (Next Line) ──
            opts = q.get("options", [])
            if q.get("type") == "MCQ" and opts and isinstance(opts, list):
                prefixes = ['a)', 'b)', 'c)', 'd)']
                for i, opt in enumerate(opts[:4]):
                    lines.append(f"   {prefixes[i]} {opt}")
                lines.append("") # Blank line after MCQ
        
        lines += ["","="*65,"  — End of Question Paper —","="*65]
        return "\n".join(lines)

    @staticmethod
    def _to_html(paper):
        """User's refined HTML formatter from snippet."""
        secs = {}
        for q in paper.get("questions", []): 
            secs.setdefault(q.get("section","Section A"),[]).append(q)
        
        rows = ""
        for sec, qs in sorted(secs.items()):
            rows += f'<tr><td colspan="4" style="background:#F4F0FB;padding:8px 12px;font-weight:800;font-size:12px;color:#7209B7;">{sec.upper()} — {len(qs)} Q · {sum(q["marks"] for q in qs)} Marks</td></tr>'
            for q in qs:
                rows += f'<tr style="border-bottom:1px solid #EDE8F5;"><td style="padding:8px;font-weight:700;width:40px;">{q["no"]}</td><td style="padding:8px;">{q["q"]}</td><td style="padding:8px;width:90px;text-align:center;"><span style="background:#F4F0FB;padding:2px 7px;border-radius:7px;font-size:10px;">{q["type"]}</span></td><td style="padding:8px;width:50px;text-align:center;font-weight:700;color:#F72585;">{q["marks"]}</td></tr>'
        
        return f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>{paper.get('title', 'Paper')}</title>
        <style>body{{font-family:Arial,sans-serif;max-width:800px;margin:40px auto;}}h1,h2{{text-align:center;}}table{{width:100%;border-collapse:collapse;font-size:13px;}}th{{background:#7209B7;color:white;padding:10px;text-align:left;}}</style></head>
        <body><h2>{paper.get("institution",paper.get("inst_name",""))}</h2><h2>{paper.get('title', 'Exam')} — {paper.get('name', 'Paper')}</h2><p style="text-align:center;font-size:12px;color:#666;">{paper.get("course",paper.get("course_name",""))} | {paper.get("dept",paper.get("department",""))} | {paper.get("semester","")} | Date: {paper.get("exam_date","")} | Duration: {paper.get("duration","")} | Max Marks: {paper.get('max_marks', paper.get('total_marks', 100))}</p><table><thead><tr><th>#</th><th>Question</th><th>Type</th><th>Marks</th></tr></thead><tbody>{rows}</tbody></table><p style="text-align:center;margin-top:30px;font-size:11px;color:#999;">— End of Question Paper —</p></body></html>"""

    @staticmethod
    def export_to_html(paper, filename=None):
        if not filename:
            filename = f"paper_{paper.get('set','A')}_{datetime.now().strftime('%H%M%S')}.html"
        filepath = os.path.join(Config.OUTPUT_DIR, filename)
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(ExportHandler._to_html(paper))
        return filepath
